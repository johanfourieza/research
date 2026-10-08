"""Rebuild revision evidence from read-only sources. All writes stay in revision."""
from pathlib import Path
import csv, json, hashlib, re, collections, datetime, statistics
import xml.etree.ElementTree as ET
import openpyxl

REV=Path(__file__).resolve().parents[1]
ROOT=REV.parent
DATA=ROOT.parent/'sources/data'
OUT=REV/'outputs'; SRC=REV/'sources'; VAL=REV/'validation'
for p in (OUT,SRC,VAL):p.mkdir(exist_ok=True)
def write(name,rows,folder=OUT):
    if not rows:return
    with (folder/name).open('w',encoding='utf-8-sig',newline='') as f:
        fields=list(dict.fromkeys(k for row in rows for k in row))
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
def read(p,delimiter=','):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f,delimiter=delimiter))
def txt(e):return re.sub(r'\s+',' ',''.join(e.itertext())).strip() if e is not None else ''
manifest=[]
def fingerprint(p):manifest.append({'path':str(p.relative_to(ROOT.parent)), 'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size})

# Document-level extraction: a heading date is not assumed to be a death date.
prob=[]; full=[]
for p in sorted((DATA/'XML files').glob('MOOC8*.xml')):
    fingerprint(p)
    for d in ET.parse(p).findall('.//div[@n]'):
        h=d.find('head');date=None if h is None else h.find('.//date')
        if date is None:continue
        dv=date.get('value','')
        if not dv[:4].isdigit():continue
        y=int(dv[:4])
        if not 1695<=y<=1720:continue
        ps=[txt(x) for x in d.findall('./p')]
        names='; '.join(txt(x) for x in h.findall('.//name[@type="person"]'))
        whole=txt(d)
        row={'id':d.get('n'),'source_file':p.name,'date_value':dv,'date_text':txt(date),'year':y,'heading_names':names,
             'opening':' '.join(ps[:2]),'closing':' '.join(ps[-2:]),'enslaved_asset_rows':len(d.findall('.//row[@role="slave"]')),
             'minor_keyword':int(bool(re.search(r'minderjar|onmondig',whole,re.I))),
             'child_keyword':int(bool(re.search(r'kinderen|kinder|dogter|dochter|soon|zoon',whole,re.I))),
             'guardian_keyword':int(bool(re.search(r'voogd|voogt|voogh',whole,re.I))),
             'widow_keyword':int(bool(re.search(r'weduw|wed:e|weedewee',whole,re.I)))}
        prob.append(row)
        if y in (1713,1714):
            blocks=[txt(x) for x in d]
            full.append(f"## {row['id']} | {dv} | {names}\n\n"+'\n\n'.join(blocks))
write('probate_register_1695_1720.csv',prob)
write('probate_register_1713_1714.csv',[r for r in prob if r['year'] in (1713,1714)])
(SRC/'probate_1713_1714_fulltext.md').write_text('# MOOC8 documents dated 1713–1714\n\n'+'\n\n'.join(full),encoding='utf-8')
(SRC/'probate_screening_text.md').write_text('\n\n'.join(f"## {r['id']} | {r['date_value']} | {r['heading_names']}\n\n{r['opening']}\n\n{r['closing']}" for r in prob if r['year'] in (1713,1714)),encoding='utf-8')
counts=collections.Counter(r['year'] for r in prob)
write('probate_by_year.csv',[{'year':y,'documents':counts[y]} for y in range(1695,1721)])
write('probate_baselines.csv',[{'start':a,'end':b,'mean':statistics.mean(counts[y] for y in range(a,b+1)),
 'ratio_1713':counts[1713]/statistics.mean(counts[y] for y in range(a,b+1)),
 'ratio_1713_1714_annualised':(counts[1713]+counts[1714])/2/statistics.mean(counts[y] for y in range(a,b+1))}
 for a,b in [(1700,1712),(1708,1712),(1709,1712)]])

# Opgaaf: preserve cell missingness and non-numeric data; only explicit integer row IDs.
files={'Cape':'Early Cape District - Including Full Index for Hague & Cape Archives September 2024.xlsx',
 'StelDrak':'Stellenbosch-Drakenstein Earlier Opgaafrolle - Hague & Cape Archives - including Full Indexes June 2022.xlsx'}
syn={'settler_men':['mannen','mans'],'settler_women':['vrouwen'],'settler_sons':['zoonen','soons','soonen'],
 'settler_daughters':['dogters','dochters'],'knechts':['knegts','knechts'],'slave_men':['slaven'],
 'slave_women':['slavinnen'],'slave_boys':['jongens','jongetjes'],'slave_girls':['meijsies','meijsjes','meisjes']}
agg=[];households=[];metadata=[];issues=[];index=[]
for district,fn in files.items():
    p=DATA/fn;fingerprint(p);w=openpyxl.load_workbook(p,read_only=True,data_only=True)
    for s in w:
        if s.title.lower().startswith('index'):
            for rn,row in enumerate(s.iter_rows(max_col=8,values_only=True),1):
                if len(row)>1 and str(row[1]).isdigit() and 1700<=int(row[1])<=1720:
                    index.append({'district_workbook':district,'sheet':s.title,'row':rn,'values':json.dumps(row,ensure_ascii=False,default=str)})
        if not s.title.isdigit() or not 1700<=int(s.title)<=1720:continue
        year=int(s.title); rows=list(s.iter_rows(max_col=80,values_only=True)); header=None
        for hi,row in enumerate(rows[:8]):
            labels=[str(x).strip().lower() if x is not None else '' for x in row]
            if 'slavinnen' in labels and ('mannen' in labels or 'mans' in labels):header=(hi,labels);break
        if header is None:
            issues.append({'district':district,'year':year,'row':0,'variable':'header','value':'no compatible header'});continue
        hi,labels=header;cmap={k:next((j for j,s in enumerate(labels) if s in choices),None) for k,choices in syn.items()}
        nrcol=next((j for j,s in enumerate(labels) if s in ['nr.','nr']),0)
        metadata.append({'district':district,'year':year,'sheet_heading':' | '.join(str(x) for x in rows[0] if x is not None),
            'header_row':hi+1,'headers':json.dumps(labels,ensure_ascii=False),'khoesan_header':int(any(re.search('hottent|khoe|inboor',s) for s in labels)),
            'missing_categories':'|'.join(k for k,v in cmap.items() if v is None)})
        hh=[]
        for rn,row in enumerate(rows[hi+1:],hi+2):
            try:int(str(row[nrcol]).strip())
            except (ValueError,TypeError):continue
            r={'district':district,'year':year,'excel_row':rn,'row_id':row[nrcol],
               'head_man':row[2] or '', 'head_woman':row[3] or ''}
            for k,j in cmap.items():
                v=row[j] if j is not None else None
                r[k+'_blank']=int(v is None or v=='')
                if j is None:r[k]=None
                elif v is None or v=='':r[k]=0
                else:
                    try:r[k]=float(v)
                    except (ValueError,TypeError):
                        r[k]=None;issues.append({'district':district,'year':year,'row':rn,'variable':k,'value':str(v)})
            hh.append(r);households.append(r)
        out={'district':district,'year':year,'households':len(hh)}
        for k in syn:
            out[k]=sum(r[k] for r in hh if r[k] is not None) if cmap[k] is not None else None
            out[k+'_blanks']=sum(r[k+'_blank'] for r in hh)
        if all(out[k] is not None for k in syn):
            for group,ac,cc in [('settler',['settler_men','settler_women'],['settler_sons','settler_daughters']),('slave',['slave_men','slave_women'],['slave_boys','slave_girls'])]:
                out[group+'_adults']=sum(out[k] for k in ac);out[group+'_children']=sum(out[k] for k in cc)
                out[group+'_total']=out[group+'_adults']+out[group+'_children']
                out[group+'_children_per_household']=out[group+'_children']/len(hh)
                out[group+'_households_with_children']=sum(sum(r[k] or 0 for k in cc)>0 for r in hh)
                out[group+'_child_adult_ratio']=out[group+'_children']/out[group+'_adults'] if out[group+'_adults'] else None
        agg.append(out)
    w.close()
# All observed sheets share the expected keys; missing columns are explicit.
write('opgaaf_household_rows.csv',households);write('opgaaf_by_year.csv',agg)
write('opgaaf_sheet_metadata.csv',metadata,SRC);write('opgaaf_index_entries.csv',index,SRC)
write('opgaaf_cell_issues.csv',issues or [{'district':'','year':'','row':'','variable':'','value':'none found'}])
changes=[]
for district in files:
    lookup={r['year']:r for r in agg if r['district']==district}
    for year,r in sorted(lookup.items()):
        if year+2 not in lookup:continue
        s=lookup[year+2]
        for group in ('settler','slave'):
            if group+'_adults' not in r or group+'_adults' not in s:continue
            a,b=r[group+'_adults'],s[group+'_adults'];c,d=r[group+'_children'],s[group+'_children']
            changes.append({'district':district,'start':year,'end':year+2,'group':group,'adults_start':a,'adults_end':b,
                'children_start':c,'children_end':d,'adult_change_pct':100*(b/a-1) if a else None,
                'child_change_pct':100*(d/c-1) if c else None,'child_change_count':d-c,
                'households_start':r['households'],'households_end':s['households']})
write('opgaaf_two_year_changes.csv',changes)

# Journal coverage and reproducible retrieval; not a claim of exhaustive semantic reading.
p=DATA/'vc_datastel (v3).xlsx';fingerprint(p)
w=openpyxl.load_workbook(p,read_only=True,data_only=True); journal=[]
for r in w.worksheets[0].iter_rows(min_row=2,max_col=4,values_only=True):
    if not r[1] or not str(r[1])[:4].isdigit():continue
    y=int(str(r[1])[:4])
    if not 1700<=y<=1720:continue
    t=re.sub('<[^>]+>',' ',str(r[3] or ''));t=re.sub(r'\s+',' ',t).strip()
    journal.append({'id':r[0],'date':str(r[1]),'year':y,'text':t,'has_text':int(t.lower() not in ('','geen teks'))})
w.close()
coverage=[]
for y in range(1700,1721):
    rs=[r for r in journal if r['year']==y]
    coverage.append({'year':y,'calendar_days':(datetime.date(y+1,1,1)-datetime.date(y,1,1)).days,
      'rows':len(rs),'unique_dates':len({r['date'] for r in rs}),'with_text':sum(r['has_text'] for r in rs)})
write('journal_coverage.csv',coverage)
rx=re.compile(r'kinder.?pok|pokjes|pokken|pocken|genees|medic|hottent|wees|weduw|ceijlon|c[eey]+lon|slav|lijveijg',re.I)
hits=[r for r in journal if 1712<=r['year']<=1714 and rx.search(r['text'])]
write('journal_candidates_1712_1714.csv',hits)
(SRC/'journal_candidates_1712_1714.md').write_text('\n\n'.join(f"## {r['id']} | {r['date']}\n\n{r['text']}" for r in hits),encoding='utf-8')

# Preserve all reference and human decisions; no model labels promoted to human gold.
gold={r['id']:r for r in read(ROOT/'notes/gold_labels.csv')}
primary={r['id']:r for r in read(ROOT/'analysis/outputs/primary_labels_1700_1720.csv')}
sample=read(ROOT/'notes/validation_sample_1700_1720.tsv','\t')
humans={}
for label,path in [('karli',ROOT/'notes/validation_79_for_student_karli.xlsx'),('martie',ROOT/'validation_79_for_student.xlsx')]:
    fingerprint(path);w=openpyxl.load_workbook(path,read_only=True,data_only=True);rs=list(w['Coding'].values)
    humans[label]={r[0]:dict(zip(rs[0],r)) for r in rs[1:]};w.close()
decisions=[]; stats=[]
for r in sample:
    i=r['id'];out={'id':i,'date':r['date'],'text':r['entry']}
    for c in ('V1','V3','V4'):
        out['rule_'+c]=primary.get(i,{}).get(c,'');out['stored_'+c]=gold.get(i,{}).get('g'+c,'')
        for label,hm in humans.items():out[label+'_'+c]=hm.get(i,{}).get('g'+c)
    for label,hm in humans.items():out[label+'_who']=hm.get(i,{}).get('gWho');out[label+'_note']=hm.get(i,{}).get('notes')
    out['stored_who']=gold.get(i,{}).get('gWho','');out['author_decision']='';out['status']='not adjudicated';decisions.append(out)
write('validation_all_400.csv',decisions,VAL)
for label,hm in humans.items():
    for c in ('V1','V3','V4'):
        pairs=[(str(r['g'+c]),gold[i]['g'+c]) for i,r in hm.items() if r.get('g'+c) is not None and i in gold]
        stats.append({'reader':label,'variable':c,'answered':len(pairs),'agreements':sum(a==b for a,b in pairs),
            'disagreements':sum(a!=b for a,b in pairs),'human_positive':sum(a=='1' for a,b in pairs)})
write('validation_agreement.csv',stats)
packet=[]
for r in decisions:
    missing=r['stored_V1']=='' or (r['id'] in humans['martie'] and r['martie_V1'] is None)
    disagree=any(r[label+'_'+c] is not None and str(r[label+'_'+c])!=str(r['stored_'+c]) for label in humans for c in ('V1','V3','V4'))
    if missing or disagree or gold.get(r['id'],{}).get('gV4')=='1':packet.append(r)
write('priority_adjudication.csv',packet,VAL)
(VAL/'priority_adjudication.md').write_text('# Author adjudication packet\n\nNo proposed AI decision is treated as human validation.\n\n'+'\n\n'.join(f"## {r['id']} | {r['date']}\n\n{r['text']}\n\n"+'; '.join(f'{k}: {v}' for k,v in r.items() if k not in ['id','date','text'] and v is not None) for r in packet),encoding='utf-8')
write('input_manifest.csv',manifest,SRC)
print(json.dumps({'probate_1713':counts[1713],'probate_1714':counts[1714],'opgaaf_district_years':len(agg),'opgaaf_household_rows':len(households),'cell_issues':len(issues),'journal_candidates':len(hits),'priority_adjudication_entries':len(packet)},indent=2))
